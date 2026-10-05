"""Builds 로또번호_분석.xlsm and LottoAnalyzer.bas.

    python3 build.py

Needs openpyxl and LibreOffice (Calc) with Python UNO.
The VBA source keeps Korean text in string literals; this script rewrites
them as U("\\uXXXX") calls so the module itself is pure ASCII and does not
depend on the VBA project code page.
"""
import os
import subprocess
import sys
import time

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
OUT_DIR = os.path.dirname(HERE)
SRC = os.path.join(HERE, "LottoAnalyzer.src.vb")
XLSM = os.path.join(OUT_DIR, "로또번호_분석.xlsm")
BAS = os.path.join(OUT_DIR, "LottoAnalyzer.bas")
CLS = os.path.join(OUT_DIR, "ThisWorkbook_code.txt")
TMP_XLSX = os.path.join(HERE, "_layout.xlsx")

DATA_FOLDER = r"C:\Users\END_USER\Dropbox\MPPT\_MicroInverter\V2000\DB\MOSFET"

THISWORKBOOK_CODE = """Private Sub Workbook_Open()
    On Error Resume Next
    EnsureButtons
End Sub
"""


# --------------------------------------------------------------------------
# VBA preprocessing
# --------------------------------------------------------------------------
def esc(s):
    return "".join(c if ord(c) < 128 else "\\u%04X" % ord(c) for c in s)


def to_ascii_vba(src):
    out_lines = []
    for line in src.splitlines():
        out, i, n = [], 0, len(line)
        while i < n:
            ch = line[i]
            if ch == '"':
                j = i + 1
                while j < n:
                    if line[j] == '"':
                        if j + 1 < n and line[j + 1] == '"':
                            j += 2
                            continue
                        break
                    j += 1
                lit = line[i:j + 1]
                body = lit[1:-1]
                if any(ord(c) > 127 for c in body):
                    out.append('U("%s")' % esc(body))
                else:
                    out.append(lit)
                i = j + 1
            elif ch == "'":
                out.append(esc(line[i:]))
                break
            else:
                out.append(ch)
                i += 1
        out_lines.append("".join(out))
    code = "\n".join(out_lines) + "\n"
    assert all(ord(c) < 128 for c in code), "non-ASCII left in VBA"
    return code


# --------------------------------------------------------------------------
# workbook layout (openpyxl)
# --------------------------------------------------------------------------
TITLE = Font(bold=True, size=14)
BOLD = Font(bold=True)
HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
NOTE = Font(color="C00000", bold=True)
THIN = Side(style="thin", color="A6A6A6")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def header(ws, row, col, labels):
    for k, lab in enumerate(labels):
        c = ws.cell(row=row, column=col + k, value=lab)
        c.font = BOLD
        c.fill = HEAD_FILL
        c.border = BOX
        c.alignment = Alignment(horizontal="center")


GUIDE = [
    ("로또 6/45 과거 1등 패턴 분석기", TITLE),
    ("", None),
    ("■ 꼭 알아두실 점", BOLD),
    ("· 로또는 매 회차가 독립된 무작위 추첨입니다. 어떤 계산으로도 당첨확률을 올릴 수 없습니다.", NOTE),
    ("· 모든 조합의 1등 확률은 똑같이 1/8,145,060 입니다.", NOTE),
    ("· 이 파일의 '점수'는 '과거 1등 번호들의 통계 패턴과 얼마나 닮았는가'를 나타낼 뿐입니다.", NOTE),
    ("", None),
    ("■ 사용 순서", BOLD),
    ("1. 매크로 사용: 파일을 열 때 '콘텐츠 사용'을 누르세요.", None),
    ("   (인터넷에서 받은 파일이라 매크로가 차단되면: 파일 우클릭 → 속성 → '차단 해제' 체크 → 확인)", None),
    ("2. [설정] 시트에서 [① 당첨번호 엑셀 불러오기] 버튼 → 당첨번호 엑셀/CSV 파일 선택", None),
    ("   · 처음 열리는 폴더는 [설정] C4 칸의 경로입니다.", None),
    ("   · 동행복권 '회차별 당첨번호' 엑셀 다운로드 파일을 그대로 쓸 수 있습니다.", None),
    ("   · 직접 만든 파일은 '회차 | 번호1~6 | 보너스' 순서로, 제목 줄에 '회차'라는 글자를 넣어주세요.", None),
    ("   · [당첨번호] 시트에 직접 붙여넣어도 됩니다. (A열 회차, B~G열 번호 6개, H열 보너스)", None),
    ("3. [② 추천번호 계산] 버튼 → [추천번호] 시트에 결과, [통계] 시트에 계산 근거가 나옵니다.", None),
    ("   · 8,145,060개 조합을 전부 계산하므로 PC에 따라 수십 초 걸릴 수 있습니다. (아래 상태표시줄에 진행률 표시)", None),
    ("", None),
    ("■ 계산 방법 (고정 규칙: 같은 데이터 + 같은 설정이면 항상 같은 결과)", BOLD),
    ("과거 1등 번호에서 아래 항목들의 분포를 셉니다.", None),
    ("  ① 번호별 출현빈도 (전체 회차)      ② 번호별 출현빈도 (최근 N회, [설정] C9)", None),
    ("  ③ 6개 번호의 합계                  ④ 홀수 개수                     ⑤ 저번호(1~22) 개수", None),
    ("  ⑥ 연속번호 쌍의 수                 ⑦ 끝수(일의 자리) 종류 수       ⑧ 한 구간(1~9,10~19,…,40~45)에 몰린 최대 개수", None),
    ("  ⑨ 직전 회차와 겹치는 번호 수       ⑩ AC값(번호 간 차이의 다양성)", None),
    ("각 조합의 점수 = Σ 가중치 × ln(과거 1등에서 그 값이 나온 비율)", None),
    ("  → 과거 1등 조합들에서 자주 나온 패턴일수록 점수가 높습니다. (번호빈도는 평균 대비 비율의 ln)", None),
    ("  → 1단계: 8,145,060개 전 조합을 ①~⑨로 채점해 상위 20,000개 보관, 2단계: ⑩ AC 점수를 더해 재정렬", None),
    ("  → 추천 조합끼리 공통번호가 [설정] C8 개수를 넘지 않도록 위에서부터 고릅니다.", None),
    ("  → [설정] C10 = Y 이면 이미 1등으로 나왔던 조합은 제외합니다.", None),
    ("가중치는 [설정] C13~C22에서 바꿀 수 있습니다. (0 = 해당 항목 미반영)", None),
]

WEIGHTS = [
    ("번호 출현빈도 (전체)", "자주 나온 번호일수록 +"),
    ("번호 출현빈도 (최근 N회)", "최근 N회(C9)에서 자주 나온 번호일수록 +"),
    ("합계", "과거 1등 합계 분포 (±7 구간)"),
    ("홀짝 비율", "홀수 개수 분포"),
    ("저고 비율", "1~22 번호 개수 분포"),
    ("연속번호", "연속번호 쌍 개수 분포"),
    ("끝수", "일의 자리 종류 수 분포"),
    ("구간 쏠림", "한 구간 최대 개수 분포"),
    ("직전회차 중복", "가장 최근 회차와 겹치는 개수 분포"),
    ("AC값", "번호 간 차이 다양성 분포"),
]


def build_layout(path):
    wb = Workbook()
    g = wb.active
    g.title = "사용법"
    g.column_dimensions["A"].width = 130
    for r, (text, font) in enumerate(GUIDE, 1):
        c = g.cell(row=r, column=1, value=text)
        if font:
            c.font = font

    s = wb.create_sheet("설정")
    s.column_dimensions["A"].width = 3
    s.column_dimensions["B"].width = 30
    s.column_dimensions["C"].width = 62
    s.column_dimensions["D"].width = 44
    s.column_dimensions["E"].width = 3
    s.column_dimensions["F"].width = 32
    s["B1"] = "로또 6/45 과거 1등 패턴 분석 - 설정"
    s["B1"].font = TITLE
    s["B2"] = "※ 모든 조합의 1등 확률은 1/8,145,060로 같습니다. 점수는 과거 1등 패턴과 닮은 정도입니다."
    s["B2"].font = NOTE
    rows = [
        (4, "데이터 폴더", DATA_FOLDER, "불러오기 창이 처음 열리는 폴더"),
        (5, "마지막 불러온 파일", None, "자동 기록"),
        (7, "추천 조합 수", 10, "1 ~ 100"),
        (8, "조합 간 최대 공통번호 수", 3, "0 ~ 5 (작을수록 서로 다른 번호로 구성)"),
        (9, "최근 회차 수 (N)", 50, "최근 출현빈도 계산 범위"),
        (10, "과거 1등 조합 제외", "Y", "Y / N"),
    ]
    for r, lab, val, desc in rows:
        s.cell(row=r, column=2, value=lab).font = BOLD
        c = s.cell(row=r, column=3, value=val)
        c.border = BOX
        if r != 5:
            c.fill = INPUT_FILL
        s.cell(row=r, column=4, value=desc)
    dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=False)
    s.add_data_validation(dv)
    dv.add("C10")
    header(s, 12, 2, ["가중치 항목", "가중치", "설명"])
    for k, (lab, desc) in enumerate(WEIGHTS):
        r = 13 + k
        s.cell(row=r, column=2, value=lab).border = BOX
        c = s.cell(row=r, column=3, value=1)
        c.fill = INPUT_FILL
        c.border = BOX
        c.alignment = Alignment(horizontal="left")
        s.cell(row=r, column=4, value=desc).border = BOX
    s["F3"] = "버튼이 안 보이면 Alt+F8 → EnsureButtons 실행"
    s["F3"].font = Font(color="808080", size=9)

    d = wb.create_sheet("당첨번호")
    header(d, 1, 1, ["회차", "번호1", "번호2", "번호3", "번호4", "번호5", "번호6", "보너스"])
    d.freeze_panes = "A2"
    for col in "ABCDEFGH":
        d.column_dimensions[col].width = 9

    st = wb.create_sheet("통계")
    st["A1"] = "계산 근거 (과거 1등 번호 통계)"
    st["A1"].font = TITLE
    st["A2"] = "기준:"
    header(st, 4, 1, ["번호", "전체 출현", "회당 출현율", "최근 출현", "번호점수"])
    header(st, 4, 7, ["항목", "값", "과거 횟수", "비율", "점수(가중치×ln)"])
    for col, wd in zip("ABCDEFGHIJK", [7, 10, 11, 10, 11, 3, 20, 10, 10, 9, 15]):
        st.column_dimensions[col].width = wd
    st.freeze_panes = "A5"

    rs = wb.create_sheet("추천번호")
    rs["A1"] = "과거 1등 패턴 유사도 상위 조합"
    rs["A1"].font = TITLE
    rs["A2"] = "기준:"
    rs["A3"] = "※ 당첨 예측이 아닙니다. 모든 조합의 1등 확률은 1/8,145,060로 같고, 점수는 과거 1등 패턴과 닮은 정도입니다."
    rs["A3"].font = NOTE
    header(rs, 4, 1, ["순위", "번호1", "번호2", "번호3", "번호4", "번호5", "번호6", "점수",
                      "합계", "홀:짝", "저:고", "연속쌍", "끝수종류", "구간최대", "직전회차중복", "AC"])
    for col, wd in zip("ABCDEFGHIJKLMNOP", [6, 7, 7, 7, 7, 7, 7, 10, 7, 7, 7, 7, 9, 9, 12, 6]):
        rs.column_dimensions[col].width = wd
    rs.freeze_panes = "A5"

    wb.active = 1
    wb.save(path)


# --------------------------------------------------------------------------
# LibreOffice: add VBA project and save as .xlsm
# --------------------------------------------------------------------------
def lo_connect(port=2002):
    import uno
    proc = subprocess.Popen(
        ["soffice", "--headless", "--invisible", "--norestore", "--nologo",
         "--accept=socket,host=localhost,port=%d;urp;" % port],
        stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    local = uno.getComponentContext()
    res = local.ServiceManager.createInstanceWithContext("com.sun.star.bridge.UnoUrlResolver", local)
    for _ in range(120):
        try:
            ctx = res.resolve("uno:socket,host=localhost,port=%d;urp;StarOffice.ComponentContext" % port)
            break
        except Exception:
            time.sleep(0.5)
    else:
        raise RuntimeError("LibreOffice did not start")
    desk = ctx.ServiceManager.createInstanceWithContext("com.sun.star.frame.Desktop", ctx)
    return proc, desk


def prop(name, value):
    import uno  # noqa: F401
    from com.sun.star.beans import PropertyValue
    p = PropertyValue()
    p.Name = name
    p.Value = value
    return p


def file_url(path):
    import uno
    return uno.systemPathToFileUrl(os.path.abspath(path))


def assemble_xlsm(layout, out, module_code):
    import uno  # noqa: F401  (enables the com.sun.star imports)
    from com.sun.star.script import ModuleInfo
    from com.sun.star.script.ModuleType import DOCUMENT, NORMAL

    proc, desk = lo_connect()
    try:
        doc = desk.loadComponentFromURL(file_url(layout), "_blank", 0, (prop("Hidden", True),))
        libs = doc.BasicLibraries
        libs.VBACompatibilityMode = True
        libs.ProjectName = "VBAProject"
        if not libs.hasByName("VBAProject"):
            libs.createLibrary("VBAProject")
        lib = libs.getByName("VBAProject")

        def add(name, code, mtype):
            mi = ModuleInfo()
            mi.ModuleType = mtype
            if lib.hasModuleInfo(name):
                lib.removeModuleInfo(name)
            if lib.hasByName(name):
                lib.removeByName(name)
            lib.insertModuleInfo(name, mi)
            lib.insertByName(name, "Option VBASupport 1\n" + code)

        add("ThisWorkbook", THISWORKBOOK_CODE, DOCUMENT)
        sheets = doc.Sheets
        for i in range(sheets.Count):
            sh = sheets.getByIndex(i)
            sh.CodeName = "Sheet%d" % (i + 1)
            add(sh.CodeName, "", DOCUMENT)
        add("LottoAnalyzer", module_code, NORMAL)
        doc.CurrentController.setActiveSheet(sheets.getByName("설정"))
        doc.storeToURL(file_url(out), (prop("FilterName", "Calc MS Excel 2007 VBA XML"),))
        doc.close(True)
    finally:
        try:
            desk.terminate()
        except Exception:
            pass
        proc.wait(timeout=30)


def add_code_names(path):
    """Excel binds document modules to sheets through codeName attributes."""
    import re
    import shutil
    import zipfile

    tmp = path + ".tmp"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/workbook.xml":
                xml = data.decode("utf-8")
                if "codeName=" not in xml:
                    xml = xml.replace("<workbookPr ", '<workbookPr codeName="ThisWorkbook" ', 1)
                data = xml.encode("utf-8")
            else:
                m = re.fullmatch(r"xl/worksheets/sheet(\d+)\.xml", item.filename)
                if m:
                    xml = data.decode("utf-8")
                    cn = 'codeName="Sheet%s"' % m.group(1)
                    if "<sheetPr" in xml:
                        xml = re.sub(r"<sheetPr(?=[\s/>])", "<sheetPr " + cn, xml, count=1)
                    else:
                        xml = re.sub(r"(<worksheet[^>]*>)", r"\1<sheetPr " + cn + "/>", xml, count=1)
                    data = xml.encode("utf-8")
            zout.writestr(item, data)
    shutil.move(tmp, path)


def main():
    src = open(SRC, encoding="utf-8").read()
    code = to_ascii_vba(src)
    with open(BAS, "w", encoding="ascii", newline="\r\n") as f:
        f.write('Attribute VB_Name = "LottoAnalyzer"\n' + code)
    with open(CLS, "w", encoding="ascii", newline="\r\n") as f:
        f.write(THISWORKBOOK_CODE)
    build_layout(TMP_XLSX)
    assemble_xlsm(TMP_XLSX, XLSM, code)
    add_code_names(XLSM)
    os.remove(TMP_XLSX)
    print("built", XLSM)
    print("built", BAS)


if __name__ == "__main__":
    sys.exit(main())
