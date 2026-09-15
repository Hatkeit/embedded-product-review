"""params.xlsx 템플릿 생성. 이 스크립트는 한 번 실행해 파일을 만들기 위한 것이다."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ARIAL = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F3864")
IN_FILL  = PatternFill("solid", fgColor="FFF7D6")   # 입력 칸
NOTE_FILL = PatternFill("solid", fgColor="EDF2F7")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# (name, value, type, min, max, unit, description)
PARAMS = [
    ("CANFD_NODE_1",            1, "uint32_t",    0, 2047, "ID",   "노드 1 메시지 식별자"),
    ("CANFD_NODE_2",            2, "uint32_t",    0, 2047, "ID",   "노드 2 메시지 식별자"),
    ("USE_CANFD_NODE",          2, "uint32_t",    1,    2, "",     "이 보드가 쓸 노드 번호 (1 또는 2)"),
    ("CANFD_HW_CHANNEL",        1, "uint8_t",     0,    1, "ch",   "CAN FD 채널 번호"),
    ("CANFD_BUFFER_INDEX",      0, "uint8_t",     0,    7, "",     "송신 버퍼 인덱스"),
    ("CANFD_DLC",               8, "uint8_t",     1,    8, "byte", "최대 수신 데이터 길이"),
    ("GPIO_INTERRUPT_PRIORITY", 7, "uint8_t",     0,    7, "",     "버튼 GPIO 인터럽트 우선순위"),
    ("#EXAMPLE_GAIN",         1.5, "float",       0,   10, "",     "예시 — 이름 앞 #는 생성에서 제외됩니다"),
]

wb = Workbook()
ws = wb.active
ws.title = "params"

ws["A1"] = "CAN FD 예제 파라미터"
ws["A1"].font = Font(name=ARIAL, size=14, bold=True, color="1F3864")
ws.merge_cells("A1:H1")

ws["A2"] = ("노란색 value 열만 수정하십시오. 나머지 열은 생성기가 읽는 메타데이터입니다. "
            "저장 후 gen_params.py 를 실행하면 gen_params.h 가 다시 만들어집니다.")
ws["A2"].font = Font(name=ARIAL, size=9, italic=True, color="595959")
ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
ws.merge_cells("A2:H2")
ws.row_dimensions[2].height = 28

HEAD = ["name", "value", "type", "min", "max", "unit", "description", "상태"]
HROW = 4
for col, text in enumerate(HEAD, start=1):
    c = ws.cell(row=HROW, column=col, value=text)
    c.font = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
    c.fill = HDR_FILL
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = BOX
ws.row_dimensions[HROW].height = 20

for i, (name, value, ctype, lo, hi, unit, desc) in enumerate(PARAMS):
    r = HROW + 1 + i
    ws.cell(row=r, column=1, value=name).font = Font(name="Consolas", size=10)
    v = ws.cell(row=r, column=2, value=value)
    v.font = Font(name="Consolas", size=10, bold=True, color="0000FF")   # 입력값 = 파랑
    v.fill = IN_FILL
    v.alignment = Alignment(horizontal="center")
    ws.cell(row=r, column=3, value=ctype).font = Font(name="Consolas", size=10)
    ws.cell(row=r, column=4, value=lo).font = Font(name=ARIAL, size=10)
    ws.cell(row=r, column=5, value=hi).font = Font(name=ARIAL, size=10)
    ws.cell(row=r, column=6, value=unit).font = Font(name=ARIAL, size=10)
    ws.cell(row=r, column=7, value=desc).font = Font(name=ARIAL, size=10)
    s = ws.cell(row=r, column=8,
                value=f'=IF(AND(B{r}>=D{r},B{r}<=E{r}),"OK","범위 초과")')
    s.font = Font(name=ARIAL, size=10, bold=True)
    s.alignment = Alignment(horizontal="center")
    for col in range(1, 9):
        ws.cell(row=r, column=col).border = BOX
    for col in (4, 5, 8):
        ws.cell(row=r, column=col).alignment = Alignment(horizontal="center")

note_r = HROW + len(PARAMS) + 2
ws.cell(row=note_r, column=1,
        value=("지원 type : uint8_t · uint16_t · uint32_t · int8_t · int16_t · int32_t · float · double · bool     |     "
               "이름 앞에 # 를 붙이면 그 행은 생성에서 제외됩니다     |     "
               "min·max 를 벗어나면 헤더를 만들지 않고 빌드가 멈춥니다"))
ws.cell(row=note_r, column=1).font = Font(name=ARIAL, size=9, color="595959")
ws.cell(row=note_r, column=1).fill = NOTE_FILL
ws.merge_cells(start_row=note_r, start_column=1, end_row=note_r, end_column=8)

for col, w in zip("ABCDEFGH", (28, 10, 11, 8, 8, 8, 40, 11)):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A5"

# ---- 사용법 시트 ----
ws2 = wb.create_sheet("사용법")
STEPS = [
    ("1. 값 수정", "params 시트의 노란색 value 열만 고칩니다. 상태 열이 '범위 초과'면 생성이 실패합니다."),
    ("2. 저장", "엑셀을 저장하고 닫습니다. 열어둔 채로 생성하면 읽기 오류가 날 수 있습니다."),
    ("3. 헤더 생성", "python gen_params.py params.xlsx <프로젝트>/generated"),
    ("4. 코드에서 사용", '#include "gen_params.h" 를 넣고 기존 #define 을 지웁니다.'),
    ("5. 빌드 자동화", "Makefile 의 PREBUILD 에 3번 명령을 걸면 make build 마다 자동 생성됩니다."),
    ("", ""),
    ("형상관리 규칙", "gen_params.h 도 커밋합니다. 엑셀·파이썬 없이도 빌드되어야 하기 때문입니다."),
    ("", "생성물은 절대 손으로 고치지 않습니다. 다음 생성에서 덮어써집니다."),
    ("", "엑셀은 바이너리라 git diff 가 안 보이므로, CSV 로도 내보내 함께 커밋하면 이력이 읽힙니다."),
    ("", ""),
    ("CSV 로도 됩니다", "openpyxl 설치가 어려우면 params.csv 로 저장해서 같은 명령에 넘기면 됩니다."),
    ("", "gen_params.py 는 .csv 를 표준 라이브러리만으로 읽습니다."),
]
ws2["A1"] = "사용법"
ws2["A1"].font = Font(name=ARIAL, size=14, bold=True, color="1F3864")
for i, (a, b) in enumerate(STEPS):
    r = 3 + i
    ca = ws2.cell(row=r, column=1, value=a)
    ca.font = Font(name=ARIAL, size=10, bold=bool(a))
    ca.alignment = Alignment(vertical="top")
    cb = ws2.cell(row=r, column=2, value=b)
    cb.font = Font(name=ARIAL, size=10)
    cb.alignment = Alignment(wrap_text=True, vertical="top")
ws2.column_dimensions["A"].width = 18
ws2.column_dimensions["B"].width = 82

wb.save("params.xlsx")
print("params.xlsx 생성 완료")
